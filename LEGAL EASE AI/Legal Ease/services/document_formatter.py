from io import BytesIO
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from utils.text_utils import (
	sanitize_text,
	parse_semicolon_terms,
	text_to_html,
)


def format_txt(text: str) -> bytes:
	clean_text = sanitize_text(text)
	return clean_text.encode("utf-8")


def _add_logo_to_docx(document: Document, logo_bytes: Optional[bytes]):
	if not logo_bytes:
		return

	try:
		paragraph = document.add_paragraph()
		paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
		run = paragraph.add_run()
		run.add_picture(BytesIO(logo_bytes), width=Inches(1.2))
	except Exception:
		pass


def format_docx(
	text: str,
	doc_type: str,
	brand_name: str = "LegalEase",
	logo_bytes: Optional[bytes] = None,
	terms: Optional[str] = None,
) -> bytes:
	document = Document()

	section = document.sections[0]
	section.top_margin = Inches(0.7)
	section.bottom_margin = Inches(0.7)
	section.left_margin = Inches(0.8)
	section.right_margin = Inches(0.8)

	normal_style = document.styles["Normal"]
	normal_style.font.name = "Times New Roman"
	normal_style.font.size = Pt(11)

	_add_logo_to_docx(document, logo_bytes)

	title = document.add_paragraph()
	title.alignment = WD_ALIGN_PARAGRAPH.CENTER
	title_run = title.add_run(doc_type)
	title_run.bold = True
	title_run.font.name = "Times New Roman"
	title_run.font.size = Pt(18)

	brand = document.add_paragraph()
	brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
	brand_run = brand.add_run(brand_name)
	brand_run.bold = True
	brand_run.font.size = Pt(11)

	clean_text = sanitize_text(text)
	for block in clean_text.split("\n\n"):
		block = block.strip()
		if not block:
			continue

		paragraph = document.add_paragraph()
		paragraph.paragraph_format.space_after = Pt(6)
		lines = block.splitlines()
		first_line = lines[0].strip()

		is_heading = (
			len(lines) == 1
			and len(first_line) < 100
			and (
				first_line.isupper()
				or first_line[:2].isdigit()
				or first_line.endswith(":")
			)
		)

		if is_heading:
			run = paragraph.add_run(block)
			run.bold = True
			run.font.name = "Times New Roman"
			run.font.size = Pt(12)
		else:
			for index, line in enumerate(lines):
				if index > 0:
					paragraph.add_run().add_break()
				paragraph.add_run(line)

	if terms:
		parsed_terms = parse_semicolon_terms(terms)
		if parsed_terms:
			heading = document.add_paragraph()
			heading_run = heading.add_run("Terms Summary")
			heading_run.bold = True

			table = document.add_table(rows=1, cols=2)
			table.style = "Table Grid"
			table.rows[0].cells[0].text = "No."
			table.rows[0].cells[1].text = "Term"

			for index, term in enumerate(parsed_terms, start=1):
				cells = table.add_row().cells
				cells[0].text = str(index)
				cells[1].text = term

	footer = section.footer.paragraphs[0]
	footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
	footer_run = footer.add_run(
		f"{brand_name} | AI-generated draft - legal review recommended"
	)
	footer_run.font.size = Pt(8)

	output = BytesIO()
	document.save(output)
	return output.getvalue()


class LegalEasePDF(FPDF):
	def __init__(self, brand_name: str):
		super().__init__()
		self.brand_name = brand_name

	def footer(self):
		self.set_y(-15)
		self.set_font("Helvetica", size=8)
		self.cell(
			0,
			10,
			f"{self.brand_name} | AI-generated draft - legal review recommended",
			align="C",
		)


def format_pdf(
	text: str,
	doc_type: str,
	brand_name: str = "LegalEase",
	logo_bytes: Optional[bytes] = None,
	terms: Optional[str] = None,
) -> bytes:
	pdf = LegalEasePDF(brand_name)
	pdf.set_auto_page_break(auto=True, margin=18)
	pdf.add_page()

	if logo_bytes:
		try:
			pdf.image(BytesIO(logo_bytes), x=85, y=12, w=40)
			pdf.ln(25)
		except Exception:
			pass

	pdf.set_font("Helvetica", "B", 16)
	pdf.cell(0, 10, doc_type, align="C")
	pdf.ln(8)

	pdf.set_font("Helvetica", "B", 10)
	pdf.cell(0, 7, brand_name, align="C")
	pdf.ln(10)

	clean_text = sanitize_text(text)
	for block in clean_text.split("\n\n"):
		block = block.strip()
		if not block:
			continue

		lines = block.splitlines()
		first_line = lines[0].strip()
		is_heading = (
			len(lines) == 1
			and len(first_line) < 100
			and (
				first_line.isupper()
				or first_line[:2].isdigit()
				or first_line.endswith(":")
			)
		)

		if is_heading:
			pdf.set_font("Helvetica", "B", 11)
			pdf.multi_cell(0, 7, first_line)
		else:
			pdf.set_font("Helvetica", size=10)
			pdf.multi_cell(0, 5.5, "\n".join(lines))

		pdf.ln(2)

	if terms:
		parsed_terms = parse_semicolon_terms(terms)
		if parsed_terms:
			pdf.set_font("Helvetica", "B", 11)
			pdf.cell(0, 7, "Terms Summary")
			pdf.ln(7)

			pdf.set_font("Helvetica", size=10)
			for index, term in enumerate(parsed_terms, start=1):
				pdf.multi_cell(0, 5.5, f"{index}. {term}")
				pdf.ln(1)

	return bytes(pdf.output())


def format_html_preview(text: str) -> str:
	return text_to_html(text)
